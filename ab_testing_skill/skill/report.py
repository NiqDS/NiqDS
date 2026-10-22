"""Three ways to show a `PilotAnalysis`, over one computation.

Same shape as `skill/output.py`: nothing here decides anything, it only
renders what `skill/analysis.py` worked out. Adding a fourth presentation
means adding a function here, not touching the analysis.

* `render_text`   -- the economic overview, in words. Goes in the email and
                     in `--output text`, and is what an agent relays.
* `write_excel`   -- the workbook an analyst opens, with line charts.
* `write_html`    -- a static dashboard, no backend, inline SVG, matching
                     the forms' theme.

The text renderer is a deterministic template filled from numbers. It is
not a free-form summary and must not become one: a report that says
something slightly different each time it runs cannot be compared with
last month's.
"""
from __future__ import annotations

import html
from pathlib import Path

from .analysis import PilotAnalysis
from .theme import THEME_CSS

__all__ = ["render_text", "write_excel", "write_html"]


def _rub(value: float | None) -> str:
    if value is None or value != value:
        return "—"
    return f"{value:,.0f}".replace(",", " ")


def _pct(value: float | None, digits: int = 1) -> str:
    if value is None or value != value:
        return "—"
    return f"{value:+.{digits}%}"


def _p(value: float | None) -> str:
    if value is None or value != value:
        return "—"
    return "< 0.0001" if value < 1e-4 else f"{value:.4f}"


# --- 1. the overview in words ---------------------------------------------

def render_text(analysis: PilotAnalysis) -> str:
    base, last = analysis.baseline, analysis.latest
    lines = [
        f"ПИЛОТ: {analysis.pilot_name} — статья «{analysis.article}»",
        f"Периодов в расчёте: {len(analysis.periods)} "
        f"({base.period} … {last.period})",
        "",
        "ГДЕ ГРУППЫ СЕЙЧАС",
        f"  КГ: {last.control_n} клиентов, в среднем {_rub(last.control_mean)} "
        f"(на старте {_rub(base.control_mean)})",
        f"  ЦГ: {last.target_n} клиентов, в среднем {_rub(last.target_mean)} "
        f"(на старте {_rub(base.target_mean)})",
        "",
        "ЭФФЕКТ ПИЛОТА (разница приростов, ЦГ сверх КГ)",
        f"  на клиента в месяц : {analysis.effect.describe()}",
    ]
    if analysis.effect_relative is not None:
        lines.append(f"  в процентах к базе : {_pct(analysis.effect_relative)}")
    lines.extend([
        f"  перестановочный тест: p = {_p(analysis.permutation_p)}",
        f"  накопленным итогом  : {_rub(analysis.cumulative_effect)} руб.",
    ])
    if analysis.paired is not None:
        lines.extend([
            "",
            f"ПО ПАРАМ ({analysis.pairs_used} пар; список разбивался матчингом)",
            f"  разница внутри пары : {analysis.paired.describe()}",
            f"  знаковый тест       : p = {_p(analysis.paired_sign_p)}",
        ])
        if analysis.pairs_broken:
            lines.append(
                f"  распалось пар       : {analysis.pairs_broken} "
                "(партнёр выбыл — пара из анализа исключена целиком)"
            )
        lines.append(
            "  на коротком списке с тяжёлым хвостом верить следует знаковому "
            "тесту: среднее здесь определяется одним-двумя крупными клиентами."
        )
    lines.extend([
        "",
        f"  ВЫВОД: {analysis.verdict}",
        "",
        "КАК ВЕДУТ СЕБЯ КЛИЕНТЫ ВНУТРИ ГРУПП",
        f"  выросли в ЦГ: {_share_text(analysis.target_share_up)} | "
        f"в КГ: {_share_text(analysis.control_share_up)}",
        f"  топ-3 клиента ЦГ дают {analysis.effect_concentration:.0%} всего изменения группы",
    ])
    if analysis.top_contributors:
        lines.append("  наибольший вклад:")
        for inn, delta in analysis.top_contributors[:3]:
            lines.append(f"    ИНН {inn}: {_rub(delta)} руб/мес")

    lines.extend(["", "ПОМЕСЯЧНО"])
    lines.append(f"    {'период':<12}{'КГ, ср.':>14}{'ЦГ, ср.':>14}{'разрыв':>14}")
    for point in analysis.periods:
        lines.append(
            f"    {point.period:<12}{_rub(point.control_mean):>14}"
            f"{_rub(point.target_mean):>14}{_rub(point.gap):>14}"
        )

    if analysis.warnings:
        lines.extend(["", "НА ЧТО ОБРАТИТЬ ВНИМАНИЕ"])
        lines.extend(f"  - {w}" for w in analysis.warnings)

    lines.extend([
        "",
        "КАК ЧИТАТЬ: эффект считается как разница приростов между группами "
        "(difference-in-differences).",
        "Два теста рядом не случайно: t-тест по средним чувствителен к крупным клиентам, "
        "перестановочный — нет.",
        "Расхождение между ними — это и есть ответ на вопрос, на ком держится эффект.",
    ])
    return "\n".join(lines)


def _share_text(share: float) -> str:
    return "—" if share != share else f"{share:.0%}"


# --- 2. the workbook ------------------------------------------------------

def write_excel(analysis: PilotAnalysis, pilot_folder: Path) -> Path:
    """Workbook with a dynamics sheet, a per-client sheet and the findings.

    Falls back to CSV when openpyxl is absent, the same way the rest of the
    exporter does -- the numbers matter more than the charts.
    """
    pilot_folder = Path(pilot_folder)
    try:
        from openpyxl import Workbook
        from openpyxl.chart import LineChart, Reference
    except ImportError:
        return _write_csv_fallback(analysis, pilot_folder)

    wb = Workbook()

    ws = wb.active
    ws.title = "Динамика"
    ws.append(["Период", "КГ: клиентов", "ЦГ: клиентов", "КГ: среднее",
               "ЦГ: среднее", "Разрыв", "Разрыв, %"])
    for point in analysis.periods:
        ws.append([
            point.period, point.control_n, point.target_n,
            round(point.control_mean, 2), round(point.target_mean, 2),
            round(point.gap, 2),
            round(point.relative_gap, 4) if point.relative_gap is not None else None,
        ])
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = "#,##0"
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
        for cell in row:
            cell.number_format = "0.0%"

    chart = LineChart()
    chart.title = f"Средний «{analysis.article}» по группам"
    chart.y_axis.title = "руб/мес на клиента"
    chart.x_axis.title = "период"
    chart.height, chart.width = 9, 20
    data = Reference(ws, min_col=4, max_col=5, min_row=1, max_row=len(analysis.periods) + 1)
    categories = Reference(ws, min_col=1, min_row=2, max_row=len(analysis.periods) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    ws.add_chart(chart, "I2")

    clients = wb.create_sheet("Клиенты")
    clients.append(["ИНН", "Группа", "Периодов", "На старте", "Сейчас",
                    "Изменение", "Изменение, %"])
    for trend in sorted(analysis.clients, key=lambda t: (t.group, -t.delta)):
        clients.append([
            trend.inn, trend.group, trend.periods,
            round(trend.first, 2), round(trend.last, 2), round(trend.delta, 2),
            round(trend.relative, 4) if trend.relative is not None else None,
        ])
    for row in clients.iter_rows(min_row=2, min_col=1, max_col=1):
        for cell in row:
            cell.number_format = "@"   # keep the leading zero on an ИНН
    for row in clients.iter_rows(min_row=2, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = "#,##0"
    for row in clients.iter_rows(min_row=2, min_col=7, max_col=7):
        for cell in row:
            cell.number_format = "0.0%"

    findings = wb.create_sheet("Выводы")
    for line in render_text(analysis).split("\n"):
        findings.append([line])
    findings.column_dimensions["A"].width = 110

    if analysis.comparability:
        checks = wb.create_sheet("Сопоставимость")
        checks.append(list(analysis.comparability[0].keys()))
        for row in analysis.comparability:
            checks.append(list(row.values()))

    for sheet in (ws, clients):
        for column, width in zip("ABCDEFG", (16, 10, 10, 16, 16, 16, 14)):
            sheet.column_dimensions[column].width = width

    path = pilot_folder / "pilot_report.xlsx"
    wb.save(path)
    return path


def _write_csv_fallback(analysis: PilotAnalysis, pilot_folder: Path) -> Path:
    import csv

    path = pilot_folder / "pilot_report.csv"
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Период", "КГ: среднее", "ЦГ: среднее", "Разрыв"])
        for point in analysis.periods:
            writer.writerow([point.period, round(point.control_mean, 2),
                             round(point.target_mean, 2), round(point.gap, 2)])
    (pilot_folder / "pilot_report.txt").write_text(render_text(analysis), encoding="utf-8")
    return path


# --- 3. the dashboard -----------------------------------------------------

def write_html(analysis: PilotAnalysis, pilot_folder: Path) -> Path:
    """Static dashboard: no backend, no fetch, inline SVG.

    Regenerated on each recalculation, so the link always shows the latest
    state. Same theme tokens as the forms so it does not look like a
    different product.
    """
    path = Path(pilot_folder) / "pilot_report.html"
    path.write_text(_html_document(analysis), encoding="utf-8")
    return path


def _svg_chart(analysis: PilotAnalysis) -> str:
    points = analysis.periods
    if len(points) < 2:
        return "<p class='muted'>Для графика нужно минимум два периода.</p>"

    width, height = 760, 280
    pad_l, pad_r, pad_t, pad_b = 70, 20, 20, 44
    values = [p.control_mean for p in points] + [p.target_mean for p in points]
    low, high = min(values), max(values)
    if high == low:
        high = low + 1.0
    span = high - low
    low -= span * 0.1
    high += span * 0.1

    def x(i: int) -> float:
        return pad_l + i * (width - pad_l - pad_r) / max(1, len(points) - 1)

    def y(value: float) -> float:
        return pad_t + (high - value) * (height - pad_t - pad_b) / (high - low)

    def polyline(series: list[float], css: str) -> str:
        pts = " ".join(f"{x(i):.1f},{y(v):.1f}" for i, v in enumerate(series))
        return f'<polyline class="{css}" points="{pts}"/>'

    grid = []
    for step in range(5):
        value = low + (high - low) * step / 4
        yy = y(value)
        grid.append(f'<line class="grid" x1="{pad_l}" y1="{yy:.1f}" x2="{width - pad_r}" y2="{yy:.1f}"/>')
        grid.append(
            f'<text class="tick" x="{pad_l - 8}" y="{yy + 4:.1f}" text-anchor="end">'
            f"{_rub(value)}</text>"
        )
    labels = []
    shown = [
        i for i in range(len(points))
        if len(points) <= 12 or i % max(1, len(points) // 10) == 0 or i == len(points) - 1
    ]
    for i in shown:
        # the end labels are anchored inward, or the first and last dates
        # run off the edge of the viewBox and get clipped
        anchor = "start" if i == 0 else "end" if i == len(points) - 1 else "middle"
        labels.append(
            f'<text class="tick" x="{x(i):.1f}" y="{height - pad_b + 20}" '
            f'text-anchor="{anchor}">{html.escape(points[i].period)}</text>'
        )
    dots = []
    for i, p in enumerate(points):
        dots.append(f'<circle class="dot cg" cx="{x(i):.1f}" cy="{y(p.control_mean):.1f}" r="3"/>')
        dots.append(f'<circle class="dot tg" cx="{x(i):.1f}" cy="{y(p.target_mean):.1f}" r="3"/>')

    return f"""<svg viewBox="0 0 {width} {height}" role="img"
     aria-label="Средний {html.escape(analysis.article)} по группам по периодам">
  {''.join(grid)}
  {polyline([p.control_mean for p in points], 'line cg')}
  {polyline([p.target_mean for p in points], 'line tg')}
  {''.join(dots)}
  {''.join(labels)}
</svg>"""


def _html_document(analysis: PilotAnalysis) -> str:
    base, last = analysis.baseline, analysis.latest
    effect_class = "ok" if analysis.effect.mean > 0 else "danger"

    warning_block = ""
    if analysis.warnings:
        items = "".join(f"<li>{html.escape(w)}</li>" for w in analysis.warnings)
        warning_block = f'<section class="panel warn"><h2>На что обратить внимание</h2><ul>{items}</ul></section>'

    period_rows = "".join(
        f"<tr><td>{html.escape(p.period)}</td><td class='num'>{p.control_n}</td>"
        f"<td class='num'>{p.target_n}</td><td class='num'>{_rub(p.control_mean)}</td>"
        f"<td class='num'>{_rub(p.target_mean)}</td><td class='num'>{_rub(p.gap)}</td></tr>"
        for p in analysis.periods
    )
    contributor_rows = "".join(
        f"<tr><td>{html.escape(inn)}</td><td class='num'>{_rub(delta)}</td></tr>"
        for inn, delta in analysis.top_contributors
    )

    return f"""<!doctype html>
<html lang="ru">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Пилот {html.escape(analysis.pilot_name)}</title>
<style>
{THEME_CSS}

  body {{ margin: 0; background: var(--bg); color: var(--text);
         font: 14px/1.5 -apple-system, "Segoe UI", Roboto, Arial, sans-serif; }}
  .wrap {{ max-width: 1000px; margin: 0 auto; padding: 24px 16px 64px; }}
  h1 {{ font-size: 22px; margin: 0 0 4px; }}
  h2 {{ font-size: 15px; margin: 0 0 12px; }}
  .sub {{ color: var(--muted); font-size: 13px; margin-bottom: 20px; }}
  .panel {{ background: var(--panel); border: 1px solid var(--border);
            border-radius: 12px; padding: 16px 18px; margin-bottom: 16px; }}
  .tiles {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(190px, 1fr)); gap: 12px; }}
  .tile .label {{ color: var(--muted); font-size: 12px; }}
  .tile .value {{ font-size: 22px; font-weight: 600; margin-top: 2px; }}
  .ok {{ color: var(--ok-text); }}
  .danger {{ color: var(--danger); }}
  .muted {{ color: var(--muted); }}
  table {{ width: 100%; border-collapse: collapse; font-size: 13px; }}
  th, td {{ text-align: left; padding: 6px 8px; border-bottom: 1px solid var(--border); }}
  th {{ color: var(--muted); font-weight: 500; }}
  td.num {{ text-align: right; font-variant-numeric: tabular-nums; }}
  .warn {{ border-color: var(--danger); }}
  .warn li {{ margin: 4px 0; }}
  svg {{ width: 100%; height: auto; }}
  .line {{ fill: none; stroke-width: 2; }}
  .line.cg {{ stroke: var(--muted); stroke-dasharray: 5 4; }}
  .line.tg {{ stroke: var(--accent); }}
  .dot.cg {{ fill: var(--muted); }}
  .dot.tg {{ fill: var(--accent); }}
  .grid {{ stroke: var(--border); stroke-width: 1; }}
  .tick {{ fill: var(--muted); font-size: 11px; }}
  .legend span {{ margin-right: 16px; font-size: 12px; color: var(--muted); }}
  .swatch {{ display: inline-block; width: 18px; height: 3px; vertical-align: middle; margin-right: 6px; }}
</style>
</head>
<body>
<div class="wrap">
  <h1>Пилот «{html.escape(analysis.pilot_name)}»</h1>
  <div class="sub">Статья «{html.escape(analysis.article)}» ·
    {len(analysis.periods)} период(ов): {html.escape(base.period)} … {html.escape(last.period)}</div>

  <section class="panel">
    <div class="tiles">
      <div class="tile"><div class="label">Эффект, руб/мес на клиента</div>
        <div class="value {effect_class}">{_rub(analysis.effect.mean)}</div></div>
      <div class="tile"><div class="label">В процентах к базе</div>
        <div class="value">{_pct(analysis.effect_relative)}</div></div>
      <div class="tile"><div class="label">Накопленным итогом</div>
        <div class="value">{_rub(analysis.cumulative_effect)}</div></div>
      <div class="tile"><div class="label">p-value (t / перестановочный)</div>
        <div class="value">{_p(analysis.effect.p_value)} / {_p(analysis.permutation_p)}</div></div>
    </div>
    <p style="margin:14px 0 0"><strong>Вывод:</strong> {html.escape(analysis.verdict)}</p>
  </section>

  {warning_block}

  <section class="panel">
    <h2>Средний «{html.escape(analysis.article)}» по группам</h2>
    {_svg_chart(analysis)}
    <div class="legend">
      <span><i class="swatch" style="background:var(--muted)"></i>КГ (контроль)</span>
      <span><i class="swatch" style="background:var(--accent)"></i>ЦГ (целевая)</span>
    </div>
  </section>

  <section class="panel">
    <h2>Помесячно</h2>
    <table>
      <thead><tr><th>Период</th><th class="num">КГ, клиентов</th><th class="num">ЦГ, клиентов</th>
        <th class="num">КГ, среднее</th><th class="num">ЦГ, среднее</th><th class="num">Разрыв</th></tr></thead>
      <tbody>{period_rows}</tbody>
    </table>
  </section>

  <section class="panel">
    <h2>Поведение внутри групп</h2>
    <p>Выросли: в ЦГ {_share_text(analysis.target_share_up)}, в КГ {_share_text(analysis.control_share_up)}.
       Топ-3 клиента ЦГ дают {analysis.effect_concentration:.0%} всего изменения группы.</p>
    <table>
      <thead><tr><th>Наибольший вклад, ИНН</th><th class="num">Изменение, руб/мес</th></tr></thead>
      <tbody>{contributor_rows}</tbody>
    </table>
  </section>

  <p class="muted">Эффект — разница приростов между группами (difference-in-differences).
  Два теста рядом не случайно: t-тест по средним чувствителен к крупным клиентам,
  перестановочный — нет; расхождение между ними показывает, на ком держится эффект.</p>
</div>
</body>
</html>"""
